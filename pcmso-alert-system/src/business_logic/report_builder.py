"""
Gerador de relatórios mensais em PDF e Excel.
Baseado na skill automated-reporting: ReportLab (PDF) + openpyxl (Excel).
"""
from __future__ import annotations

import logging
from datetime import date
from pathlib import Path
from typing import Optional

from sqlalchemy import and_, func, select
from sqlalchemy.orm import Session

from src.database.models import AlertaEnviado, Colaborador, Empresa, Exame, StatusExame, StatusEnvio

logger = logging.getLogger(__name__)


def _get_dados_empresa_mes(db: Session, empresa_id: int, ano: int, mes: int) -> dict:
    """Coleta métricas mensais de uma empresa."""
    from calendar import monthrange
    from datetime import date as dt

    primeiro_dia = dt(ano, mes, 1)
    ultimo_dia   = dt(ano, mes, monthrange(ano, mes)[1])

    total_col = db.scalar(select(func.count(Colaborador.id)).where(
        and_(Colaborador.empresa_id == empresa_id, Colaborador.ativo == True)
    )) or 0

    realizados = db.scalar(select(func.count(Exame.id)).join(Colaborador).where(
        and_(
            Colaborador.empresa_id == empresa_id,
            Exame.status == StatusExame.REALIZADO,
            Exame.updated_at >= primeiro_dia,
            Exame.updated_at <= ultimo_dia,
        )
    )) or 0

    pendentes = db.scalar(select(func.count(Exame.id)).join(Colaborador).where(
        and_(Colaborador.empresa_id == empresa_id, Exame.status == StatusExame.PENDENTE)
    )) or 0

    vencidos = db.scalar(select(func.count(Exame.id)).join(Colaborador).where(
        and_(Colaborador.empresa_id == empresa_id, Exame.status == StatusExame.VENCIDO)
    )) or 0

    alertas = db.scalar(select(func.count(AlertaEnviado.id)).where(
        and_(
            AlertaEnviado.empresa_id == empresa_id,
            AlertaEnviado.status_envio == StatusEnvio.ENVIADO,
            func.date(AlertaEnviado.data_envio) >= primeiro_dia,
            func.date(AlertaEnviado.data_envio) <= ultimo_dia,
        )
    )) or 0

    return {
        "total_colaboradores": total_col,
        "realizados": realizados,
        "pendentes": pendentes,
        "vencidos": vencidos,
        "alertas_enviados": alertas,
    }


def gerar_relatorio_mensal(db: Session, ano: Optional[int] = None, mes: Optional[int] = None) -> dict:
    """
    Gera relatório mensal para todas as empresas ativas.
    Envia por email e retorna resumo.
    """
    from src.database.queries import listar_empresas_ativas
    from src.communications.gmail_sender import enviar_email
    from src.business_logic.message_generator import gerar_email_relatorio_mensal
    import calendar

    hoje = date.today()
    ano  = ano or hoje.year
    mes  = mes or hoje.month
    mes_ref = f"{calendar.month_name[mes].capitalize()}/{ano}"

    empresas = listar_empresas_ativas(db)
    resultados = []

    for empresa in empresas:
        dados = _get_dados_empresa_mes(db, empresa.id, ano, mes)
        dados["mes_referencia"] = mes_ref

        # Gerar PDF e Excel
        pdf_path  = gerar_pdf_empresa(empresa, dados, ano, mes)
        xlsx_path = gerar_excel_empresa(empresa, dados, ano, mes)

        # Enviar por email
        html = gerar_email_relatorio_mensal(empresa, dados)
        assunto = f"📊 Relatório Mensal PCMSO — {empresa.nome_fantasia or empresa.razao_social} — {mes_ref}"
        sucesso, erro = enviar_email(empresa.email_sso, assunto, html, anexos=[pdf_path, xlsx_path])

        resultados.append({
            "empresa": empresa.razao_social,
            "enviado": sucesso,
            "erro": erro,
            "dados": dados,
        })
        logger.info(f"Relatório mensal {'enviado' if sucesso else 'FALHOU'}: {empresa.razao_social}")

    return {"ok": True, "mes": mes_ref, "empresas_processadas": len(resultados), "resultados": resultados}


def gerar_pdf_empresa(empresa: Empresa, dados: dict, ano: int, mes: int) -> Optional[str]:
    """Gera PDF do relatório mensal usando ReportLab."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import (
            PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
        )
        import calendar

        output_dir = Path("data/relatorios")
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / f"relatorio_{empresa.cnpj.replace('.','').replace('/','').replace('-','')}_{ano}{mes:02d}.pdf"

        doc = SimpleDocTemplate(str(output_path), pagesize=A4)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            "Title", parent=styles["Heading1"],
            fontSize=18, textColor=colors.HexColor("#2c3e50"), spaceAfter=6,
        )
        sub_style = ParagraphStyle(
            "Sub", parent=styles["Normal"],
            fontSize=11, textColor=colors.HexColor("#7f8c8d"), spaceAfter=20,
        )

        story.append(Paragraph("Relatório Mensal PCMSO", title_style))
        nome = empresa.nome_fantasia or empresa.razao_social
        story.append(Paragraph(f"{nome} — {calendar.month_name[mes].capitalize()}/{ano}", sub_style))
        story.append(Spacer(1, 0.2 * inch))

        # Tabela de métricas
        table_data = [
            ["Métrica", "Valor"],
            ["Total de Colaboradores Ativos", str(dados.get("total_colaboradores", 0))],
            ["Exames Realizados no Mês", str(dados.get("realizados", 0))],
            ["Exames Pendentes", str(dados.get("pendentes", 0))],
            ["Exames Vencidos", str(dados.get("vencidos", 0))],
            ["Alertas Enviados", str(dados.get("alertas_enviados", 0))],
        ]
        t = Table(table_data, colWidths=[4 * inch, 2 * inch])
        t.setStyle(TableStyle([
            ("BACKGROUND",  (0, 0), (-1, 0), colors.HexColor("#2c3e50")),
            ("TEXTCOLOR",   (0, 0), (-1, 0), colors.white),
            ("FONTNAME",    (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f9fa")]),
            ("GRID",        (0, 0), (-1, -1), 0.5, colors.HexColor("#dee2e6")),
            ("PADDING",     (0, 0), (-1, -1), 8),
        ]))
        story.append(t)

        doc.build(story)
        logger.info(f"PDF gerado: {output_path}")
        return str(output_path)

    except Exception as e:
        logger.error(f"Erro ao gerar PDF: {e}")
        return None


def gerar_excel_empresa(empresa: Empresa, dados: dict, ano: int, mes: int) -> Optional[str]:
    """Gera Excel do relatório mensal usando openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        import calendar

        output_dir = Path("data/relatorios")
        output_dir.mkdir(parents=True, exist_ok=True)
        output_path = output_dir / f"relatorio_{empresa.cnpj.replace('.','').replace('/','').replace('-','')}_{ano}{mes:02d}.xlsx"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Resumo"

        header_font  = Font(bold=True, color="FFFFFF", size=12)
        header_fill  = PatternFill(start_color="2C3E50", fill_type="solid")
        center_align = Alignment(horizontal="center")

        # Título
        ws["A1"] = f"Relatório Mensal PCMSO — {empresa.nome_fantasia or empresa.razao_social}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"{calendar.month_name[mes].capitalize()}/{ano}"
        ws["A2"].font = Font(size=12, color="7F8C8D")
        ws.append([])

        # Headers
        headers = ["Métrica", "Valor"]
        ws.append(headers)
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        # Dados
        rows = [
            ("Total de Colaboradores Ativos", dados.get("total_colaboradores", 0)),
            ("Exames Realizados no Mês",       dados.get("realizados", 0)),
            ("Exames Pendentes",               dados.get("pendentes", 0)),
            ("Exames Vencidos",                dados.get("vencidos", 0)),
            ("Alertas Enviados",               dados.get("alertas_enviados", 0)),
        ]
        for row in rows:
            ws.append(list(row))

        ws.column_dimensions["A"].width = 40
        ws.column_dimensions["B"].width = 15

        wb.save(str(output_path))
        logger.info(f"Excel gerado: {output_path}")
        return str(output_path)

    except Exception as e:
        logger.error(f"Erro ao gerar Excel: {e}")
        return None
