"""
Gera arquivo Excel (.xlsx) com 5 abas de dados + 1 aba de erros
a partir de uma lista de resultados de extração de PCMSOs.
"""
from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Definição das abas
# ---------------------------------------------------------------------------

ABAS: dict[str, list[str]] = {
    "Empresas": [
        "arquivo", "razao_social", "cnpj", "nome_fantasia", "grau_risco",
        "cnae", "endereco", "num_profissionais", "medico_pcmso",
        "responsavel_empresa", "vigencia_inicio", "vigencia_fim",
        "email_sso", "whatsapp_sso",
    ],
    "Cargos": ["arquivo", "razao_social", "setor", "cargo", "quantidade"],
    "Riscos": ["arquivo", "razao_social", "tipo_risco", "descricao_risco", "danos_saude"],
    "Exames": ["arquivo", "razao_social", "exame", "periodicidade"],
    "Cargo_Exames": [
        "arquivo", "razao_social", "cargo", "setor", "risco",
        "tipo_exame", "periodicidade_meses",
    ],
    "Erros": ["arquivo", "erro"],
}

_COR_HEADER = "2E75B6"  # azul escuro
_COR_ERRO = "C00000"    # vermelho para aba de erros


# ---------------------------------------------------------------------------
# Helpers de formatação
# ---------------------------------------------------------------------------

def _escrever_header(ws, colunas: list[str], cor_hex: str = _COR_HEADER) -> None:
    fill = PatternFill("solid", fgColor=cor_hex)
    font = Font(bold=True, color="FFFFFF")
    for col_idx, nome in enumerate(colunas, 1):
        cell = ws.cell(row=1, column=col_idx, value=nome)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def _auto_largura(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)


def _val(obj: Any) -> str:
    """Converte valor para string limpa, None/nan → ''."""
    if obj is None:
        return ""
    s = str(obj).strip()
    return "" if s.lower() in ("none", "nan") else s


# ---------------------------------------------------------------------------
# Função principal
# ---------------------------------------------------------------------------

def gerar_excel(resultados: list[dict]) -> bytes:
    """
    Recebe lista de dicts retornados por `extrair_pcmso()` e gera um .xlsx.

    Retorna os bytes do arquivo para facilitar resposta HTTP ou gravação em disco.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove aba padrão vazia

    # Criar todas as abas
    sheets: dict[str, Any] = {}
    for nome_aba, colunas in ABAS.items():
        ws = wb.create_sheet(nome_aba)
        cor = _COR_ERRO if nome_aba == "Erros" else _COR_HEADER
        _escrever_header(ws, colunas, cor)
        sheets[nome_aba] = ws

    for extracao in resultados:
        arquivo = _val(extracao.get("arquivo"))
        empresa = extracao.get("empresa") or {}
        razao = _val(empresa.get("razao_social"))

        # --- Aba Empresas ---
        ws_emp = sheets["Empresas"]
        ws_emp.append([
            arquivo,
            razao,
            _val(empresa.get("cnpj")),
            _val(empresa.get("nome_fantasia")),
            _val(empresa.get("grau_risco")),
            _val(empresa.get("cnae")),
            _val(empresa.get("endereco")),
            _val(empresa.get("num_profissionais")),
            _val(empresa.get("medico_pcmso")),
            _val(empresa.get("responsavel_empresa")),
            _val(empresa.get("vigencia_inicio")),
            _val(empresa.get("vigencia_fim")),
            _val(empresa.get("email_sso")),
            _val(empresa.get("whatsapp_sso")),
        ])

        # --- Aba Cargos ---
        ws_cargos = sheets["Cargos"]
        for cargo in extracao.get("cargos") or []:
            ws_cargos.append([
                arquivo, razao,
                _val(cargo.get("setor")),
                _val(cargo.get("cargo")),
                _val(cargo.get("quantidade")),
            ])

        # --- Aba Riscos ---
        ws_riscos = sheets["Riscos"]
        for risco in extracao.get("riscos") or []:
            ws_riscos.append([
                arquivo, razao,
                _val(risco.get("tipo_risco")),
                _val(risco.get("descricao_risco")),
                _val(risco.get("danos_saude")),
            ])

        # --- Aba Exames ---
        ws_exames = sheets["Exames"]
        for exame in extracao.get("exames") or []:
            ws_exames.append([
                arquivo, razao,
                _val(exame.get("exame")),
                _val(exame.get("periodicidade")),
            ])

        # --- Aba Cargo_Exames ---
        ws_ce = sheets["Cargo_Exames"]
        for ce in extracao.get("cargo_exames") or []:
            ws_ce.append([
                arquivo, razao,
                _val(ce.get("cargo")),
                _val(ce.get("setor")),
                _val(ce.get("risco")),
                _val(ce.get("tipo_exame")),
                ce.get("periodicidade_meses") or "",
            ])

        # --- Aba Erros ---
        ws_erros = sheets["Erros"]
        for erro in extracao.get("erros_extracao") or []:
            ws_erros.append([arquivo, _val(erro)])

    # Ajustar largura de colunas
    for ws in sheets.values():
        _auto_largura(ws)

    # Serializar para bytes
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
