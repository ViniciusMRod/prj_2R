"""
Converte o Excel editado pelo técnico de segurança de volta para
lista de dicts no formato esperado por ValidacaoPendente.dados_extraidos.
"""
from __future__ import annotations

import io
from collections import defaultdict

import openpyxl


def _val(cell_value) -> str:
    """Normaliza valor de célula para string limpa."""
    if cell_value is None:
        return ""
    s = str(cell_value).strip()
    return "" if s.lower() in ("none", "nan") else s


def _rows_to_dicts(ws) -> list[dict]:
    """Lê uma worksheet e retorna lista de dicts usando a primeira linha como header."""
    headers = [_val(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = {headers[i]: _val(v) for i, v in enumerate(row) if i < len(headers)}
        rows.append(d)
    return rows


def converter_excel_para_registros(excel_bytes: bytes) -> list[dict]:
    """
    Lê o Excel editado pelo técnico e retorna uma lista de dicts,
    um por arquivo/empresa, no formato de ValidacaoPendente.dados_extraidos.

    Estrutura retornada por item:
    {
        "arquivo": str,
        "empresa": {...},
        "cargos": [...],
        "riscos": [...],
        "exames": [...],
        "cargo_exames": [...],
    }
    """
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)

    # Ler cada aba (ignora silenciosamente se não existir)
    def ler_aba(nome: str) -> list[dict]:
        if nome not in wb.sheetnames:
            return []
        return _rows_to_dicts(wb[nome])

    empresas_rows = ler_aba("Empresas")
    cargos_rows = ler_aba("Cargos")
    riscos_rows = ler_aba("Riscos")
    exames_rows = ler_aba("Exames")
    ce_rows = ler_aba("Cargo_Exames")
    meta_rows = ler_aba("Metadados")

    # Agrupar tudo por "arquivo"
    registros: dict[str, dict] = {}

    for row in empresas_rows:
        arquivo = row.get("arquivo", "")
        if not arquivo:
            continue
        registros[arquivo] = {
            "arquivo": arquivo,
            "empresa": {
                "razao_social": row.get("razao_social", ""),
                "cnpj": row.get("cnpj", ""),
                "nome_fantasia": row.get("nome_fantasia", ""),
                "grau_risco": row.get("grau_risco", ""),
                "cnae": row.get("cnae", ""),
                "endereco": row.get("endereco", ""),
                "num_profissionais": row.get("num_profissionais", ""),
                "medico_pcmso": row.get("medico_pcmso", ""),
                "responsavel_empresa": row.get("responsavel_empresa", ""),
                "vigencia_inicio": row.get("vigencia_inicio", ""),
                "vigencia_fim": row.get("vigencia_fim", ""),
                "email_sso": row.get("email_sso", ""),
                "whatsapp_sso": row.get("whatsapp_sso", ""),
            },
            "cargos": [],
            "riscos": [],
            "exames": [],
            "cargo_exames": [],
            "hash": "",
            "status_pcmso": "NOVO",
        }

    # Agregar linhas das demais abas
    for row in cargos_rows:
        arquivo = row.get("arquivo", "")
        if arquivo in registros:
            registros[arquivo]["cargos"].append({
                "setor": row.get("setor", ""),
                "cargo": row.get("cargo", ""),
                "quantidade": row.get("quantidade", ""),
            })

    for row in riscos_rows:
        arquivo = row.get("arquivo", "")
        if arquivo in registros:
            registros[arquivo]["riscos"].append({
                "tipo_risco": row.get("tipo_risco", ""),
                "descricao_risco": row.get("descricao_risco", ""),
                "danos_saude": row.get("danos_saude", ""),
            })

    for row in exames_rows:
        arquivo = row.get("arquivo", "")
        if arquivo in registros:
            registros[arquivo]["exames"].append({
                "exame": row.get("exame", ""),
                "periodicidade": row.get("periodicidade", ""),
            })

    for row in ce_rows:
        arquivo = row.get("arquivo", "")
        if arquivo in registros:
            period = row.get("periodicidade_meses", "")
            registros[arquivo]["cargo_exames"].append({
                "cargo": row.get("cargo", ""),
                "setor": row.get("setor", ""),
                "risco": row.get("risco", ""),
                "tipo_exame": row.get("tipo_exame", ""),
                "periodicidade_meses": int(period) if str(period).isdigit() else None,
            })

    # Metadados: hash + status de versionamento (preserva o hash no round-trip,
    # essencial para registrar_versao na aprovação)
    for row in meta_rows:
        arquivo = row.get("arquivo", "")
        if arquivo in registros:
            registros[arquivo]["hash"] = row.get("hash", "")
            registros[arquivo]["status_pcmso"] = row.get("status_pcmso", "") or "NOVO"

    return list(registros.values())
